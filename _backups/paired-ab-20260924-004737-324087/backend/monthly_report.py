"""Two editable Word reports based on mailbox receipt month and declared plans."""
import io
import re
from datetime import datetime, timezone, timedelta
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from mmlab_pipeline.members import fold
from monthly_assembly import assemble, TOPICS

VN=timezone(timedelta(hours=7))
SECTIONS=('a1','a2','aOther','b1','b2','bOther','recommendations')
DEFAULT_STRATEGY='Chiến lược Phát triển Trường, giai đoạn 2021-2030'
PLAN_REFERENCE='Kế hoạch hành động năm 2026: https://link.uit.edu.vn/KH2026'


def validate_period(period):
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',period):
        raise ValueError('Kỳ phải có dạng YYYY-MM.')
    datetime.strptime(period,'%Y-%m')
    return period


def next_period(period):
    y,m=map(int,validate_period(period).split('-'))
    return f'{y+int(m==12):04d}-{1 if m==12 else m+1:02d}'


def work_key(r):
    return (r['type'],fold(r['title']).strip(),re.sub(r'\s+','',fold(r.get('venue',''))))


def useful(lines):
    return list(dict.fromkeys(s.strip() for s in lines if s.strip() and re.search(r'\w',s)))


def compose(records, period):
    validate_period(period)
    return assemble(records,period,SECTIONS,DEFAULT_STRATEGY,work_key)


def docx_bytes(report, variant):
    if variant not in ('discussion','school'):raise ValueError('Unknown report variant')
    # Work on a copy: rendering must not mutate the admin's saved draft.
    report={**report,'sections':dict(report['sections'])}
    if report['strategyLabel']=='KHCL Trường giai đoạn 2021-2025':report['strategyLabel']=DEFAULT_STRATEGY
    if variant=='school':
        # Only explicitly tagged strategic tasks enter A.1/B.1.
        # Other tasks are regular work; preserve their source text.
        for phase in ('a','b'):
            report['sections'][phase+'2']='\n'.join(filter(None,[report['sections'].get(phase+'2','').strip(),report['sections'].get(phase+'Other','').strip()]))
            report['sections'][phase+'Other']=''
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.line_spacing=1.08
    for style_name in ('Heading 1','Heading 2'):
        style=doc.styles[style_name];style.font.name='Times New Roman';style.font.size=Pt(12);style.font.bold=True
        style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(8)
        style.paragraph_format.space_after=Pt(4)
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            fonts.attrib.pop(qn('w:'+attr),None)
    def paragraph(text='',bold=False,center=False):
        p=doc.add_paragraph();p.paragraph_format.keep_together=True
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(text).bold=bold
        return p
    paragraph('TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN',True,True)
    paragraph('ĐƠN VỊ: PHÒNG THÍ NGHIỆM TRUYỀN THÔNG ĐA PHƯƠNG TIỆN',True,True)
    if variant=='school':paragraph('Mã đơn vị: 6',True)
    template=report.get('isTemplate',False)
    stamp=datetime.fromisoformat(report['generatedAt']).astimezone(VN)
    p=paragraph('Ngày … tháng … năm ……' if template else f'Ngày {stamp.day:02d} tháng {stamp.month:02d} năm {stamp.year}');p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    y,m=map(int,report['planPeriod'].split('-'))
    paragraph('BÁO CÁO CÔNG TÁC'+(' THÁNG …/……' if template else f' THÁNG {m:02d}/{y}'),True,True)
    if not template:
        paragraph(f"Kết quả tháng {report['period'][5:]}/{report['period'][:4]} và kế hoạch tháng {m:02d}/{y}",center=True)
    def section(key,title):
        if variant!='school' and not template and not report['sections'].get(key,'').strip():return
        if template and key=='b1' and variant=='discussion':doc.add_page_break()
        doc.add_heading(title,level=1)
        if key in ('a1','b1','b2'):
            paragraph(PLAN_REFERENCE)
        text=report['sections'].get(key,'').strip()
        for line in text.splitlines() if text else ['…' if template else 'Chưa có nội dung được ghi nhận.']:
            if line.rstrip(':') in TOPICS:
                doc.add_heading(line.rstrip(':'),level=2)
            else:
                p=paragraph(line)
                if not template:
                    p.paragraph_format.left_indent=Cm(.35)
                    p.paragraph_format.first_line_indent=Cm(-.35)
                    p.runs[0].text='• '+line
    if variant in ('discussion','school'):
        section('a1',f"Phần A.1: Tình hình thực hiện nhiệm vụ phục vụ {report['strategyLabel']} trong tháng trước")
        section('a2','Phần A.2: Tình hình thực hiện nhiệm vụ thường xuyên và đột xuất trong tháng trước')
        if report['sections'].get('aOther'):section('aOther', 'Phần A: Kết quả công tác trong tháng trước' if not report['sections'].get('a1') and not report['sections'].get('a2') else 'Nội dung thực hiện bổ sung')
    section('b1',f"Phần B.1: Kế hoạch công tác các nhiệm vụ phục vụ {report['strategyLabel']} trong tháng này")
    section('b2','Phần B.2: Kế hoạch công tác các nhiệm vụ thường xuyên và đột xuất trong tháng này')
    if report['sections'].get('bOther'):section('bOther','Phần B: Kế hoạch công tác trong tháng này' if not report['sections'].get('b1') and not report['sections'].get('b2') else 'Nhiệm vụ dự kiến bổ sung')
    if variant=='discussion' and report.get('reviewNotes'):
        doc.add_heading('Nội dung cần bổ sung để hoàn thiện báo cáo',level=1)
        for note in report['reviewNotes']:paragraph(note)
    if variant=='school' or template or report['sections'].get('recommendations','').strip():
        section('recommendations','Phần 3: Các kiến nghị')
    if doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
    p=paragraph('Trưởng đơn vị\n(Ký và ghi rõ họ tên)\n\n'+report['signatory'],True)
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    doc.core_properties.title=('Báo cáo thảo luận ' if variant=='discussion' else 'Báo cáo nộp trường ')+report['planPeriod']
    doc.core_properties.author='MMLab - UIT'
    out=io.BytesIO();doc.save(out);return out.getvalue()


def blank_template(variant):
    report=compose([], '2026-09')
    report['isTemplate']=True
    done=['Được chấp nhận … bài báo hội nghị rank A*/A/B/C/Scopus',
          'Được chấp nhận … bài báo tạp chí ISI/Scopus Q1/Q2/Q3/Q4',
          'Đã nộp … bài báo hội nghị rank A*/A/B/C/Scopus',
          'Đã nộp … bài báo tạp chí ISI/Scopus Q1/Q2/Q3/Q4',
          'Hoàn tất đăng ký … đề tài NCKH cấp cơ sở/cấp ĐHQG/NAFOSTED',
          'Hoàn tất nghiệm thu … đề tài NCKH cấp cơ sở/cấp ĐHQG/NAFOSTED',
          '… NCS đã hoàn thành nhập học/báo cáo CĐ 1/CĐ 2/CĐ 3/TLTQ/Seminar/ĐVCM/Cấp Trường',
          'Đã thực hiện báo cáo … seminar học thuật tại PTN', 'Đã đạt thành tích giải thưởng …']
    plans=['Sẽ nộp … bài báo hội nghị rank A*/A/B/C/Scopus',
           'Sẽ nộp … bài báo tạp chí ISI/Scopus Q1/Q2/Q3/Q4',
           'Sẽ đăng ký … đề tài NCKH cấp cơ sở/cấp ĐHQG/NAFOSTED',
           'Sẽ nghiệm thu … đề tài NCKH cấp cơ sở/cấp ĐHQG/NAFOSTED',
           '… NCS sẽ nhập học/báo cáo CĐ 1/CĐ 2/CĐ 3/TLTQ/Seminar/ĐVCM/Cấp Trường',
           'Sẽ thực hiện … seminar học thuật tại PTN', 'Dự kiến tham gia giải thưởng/cuộc thi …']
    report['sections'].update(a1='\n'.join(done),a2='\n'.join(done),b1='\n'.join(plans),b2='\n'.join(plans),recommendations='…')
    return docx_bytes(report,variant)
