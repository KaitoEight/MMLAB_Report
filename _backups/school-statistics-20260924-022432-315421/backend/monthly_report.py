"""Two editable Word reports based on mailbox receipt month and declared plans."""
import io
import re
from datetime import datetime, timezone, timedelta
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from mmlab_pipeline.members import fold
from monthly_assembly import assemble, TOPICS
from monthly_layout import paired_report, LAYOUT
from school_report import school_projection

VN=timezone(timedelta(hours=7))
SECTIONS=('a1','a2','aOther','b1','b2','bOther','recommendations')
DEFAULT_STRATEGY='Chiến lược Phát triển Trường, giai đoạn 2021-2030'
PLAN_REFERENCE='Kế hoạch hành động năm 2026: https://link.uit.edu.vn/KH2026'


def validate_period(period):
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',period):
        raise ValueError('Kỳ phải có dạng YYYY-MM.')
    datetime.strptime(period,'%Y-%m')
    return period


def next_period(period):
    y,m=map(int,validate_period(period).split('-'))
    return f'{y+int(m==12):04d}-{1 if m==12 else m+1:02d}'


def work_key(r):
    return (r['type'],fold(r['title']).strip(),re.sub(r'\s+','',fold(r.get('venue',''))))


def useful(lines):
    return list(dict.fromkeys(s.strip() for s in lines if s.strip() and re.search(r'\w',s)))


def compose(records, period):
    validate_period(period)
    return paired_report({**assemble(records,period,SECTIONS,DEFAULT_STRATEGY,work_key),'layoutVersion':LAYOUT})


def docx_bytes(report, variant):
    if variant not in ('discussion','school'):raise ValueError('Unknown report variant')
    report=school_projection(report) if variant=='school' else paired_report(report)
    if report['strategyLabel']=='KHCL Trường giai đoạn 2021-2025':report['strategyLabel']=DEFAULT_STRATEGY
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.line_spacing=1.08
    for style_name in ('Heading 1','Heading 2'):
        style=doc.styles[style_name];style.font.name='Times New Roman';style.font.size=Pt(12);style.font.bold=True
        style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(8)
        style.paragraph_format.space_after=Pt(4)
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            fonts.attrib.pop(qn('w:'+attr),None)
    def paragraph(text='',bold=False,center=False):
        p=doc.add_paragraph();p.paragraph_format.keep_together=True
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(text).bold=bold
        return p
    paragraph('TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN',True,True)
    paragraph('ĐƠN VỊ: PHÒNG THÍ NGHIỆM TRUYỀN THÔNG ĐA PHƯƠNG TIỆN',True,True)
    if variant=='school':paragraph('Mã đơn vị: 6',True)
    template=report.get('isTemplate',False)
    stamp=datetime.fromisoformat(report['generatedAt']).astimezone(VN)
    p=paragraph('Ngày … tháng … năm ……' if template else f'Ngày {stamp.day:02d} tháng {stamp.month:02d} năm {stamp.year}');p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    y,m=map(int,report['planPeriod'].split('-'))
    paragraph('BÁO CÁO CÔNG TÁC'+(' THÁNG …/……' if template else f' THÁNG {m:02d}/{y}'),True,True)
    if not template:
        paragraph(f"Kết quả tháng {report['period'][5:]}/{report['period'][:4]} và kế hoạch tháng {m:02d}/{y}",center=True)
    def section(key,title):
        if template and key=='b1' and variant=='discussion':doc.add_page_break()
        doc.add_heading(title,level=1)
        if variant=='school':
            if key in ('a1','a2'):
                paragraph('Với mỗi nhiệm vụ, cho biết kết quả thực hiện, đánh giá chất lượng, % hoàn thành, lý do chưa hoàn thành. Xin tham khảo kế hoạch đã xây dựng trong tháng trước.')
            if key=='a1':
                paragraph('Đề nghị phần này CHỈ báo cáo những nội dung KHCL nào có trong Kế hoạch tại https://link.uit.edu.vn/KH2026, báo cáo ngắn gọn 1–2 dòng.')
            if key in ('b1','b2'):
                paragraph('Lưu ý đối chiếu với Kế hoạch Trường 2026 đã xây dựng tại https://link.uit.edu.vn/KH2026')
        text=report['sections'].get(key,'').strip()
        for line in text.splitlines() if text else ['…' if template else 'Chưa có nội dung được ghi nhận.']:
            if line.rstrip(':') in TOPICS:
                doc.add_heading(line.rstrip(':'),level=2)
            else:
                p=paragraph(line)
                if not template:
                    p.paragraph_format.left_indent=Cm(.35)
                    p.paragraph_format.first_line_indent=Cm(-.35)
                    p.runs[0].text='• '+line
    if variant=='school':
        section('a1',f"Phần A.1: Báo cáo tình hình thực hiện nhiệm vụ phục vụ {report['strategyLabel']} trong tháng trước")
        section('a2','Phần A.2: Báo cáo tình hình thực hiện nhiệm vụ thường xuyên và đột xuất khác trong tháng trước')
        section('b1',f"Phần B.1: Kế hoạch công tác các nhiệm vụ phục vụ {report['strategyLabel']} trong tháng này")
        section('b2','Phần B.2: Kế hoạch công tác các nhiệm vụ thường xuyên và đột xuất trong tháng này')
    else:
        section('a1','Phần A: Công việc đã thực hiện')
        section('b1','Phần B: Kế hoạch công việc')
    if doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
    p=paragraph('Trưởng đơn vị\n(Ký và ghi rõ họ tên)\n\n'+report['signatory'],True)
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    doc.core_properties.title=('Báo cáo thảo luận ' if variant=='discussion' else 'Báo cáo nộp trường ')+report['planPeriod']
    doc.core_properties.author='MMLab - UIT'
    out=io.BytesIO();doc.save(out);return out.getvalue()


def blank_template(variant):
    report=compose([], '2026-09')
    report['isTemplate']=True
    done=['KPI bài báo: …','KPI NCS: …','Đề tài NCKH: …']
    plans=['KPI bài báo: …','KPI NCS: …','Đề tài NCKH: …']
    report['sections'].update(a1='\n'.join(done),a2='\n'.join(done),b1='\n'.join(plans),b2='\n'.join(plans),recommendations='…')
    return docx_bytes(report,variant)
