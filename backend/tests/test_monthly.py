import base64
import copy
import io
from datetime import datetime, timezone
from unittest.mock import patch
from test_api import client, login, bundle
import pytest
from docx import Document
from sqlalchemy import select, update
import mailer
import monthly_service as service
import store
from monthly_report import VN, compose, docx_bytes, blank_template
from report_cleanup import clean_report
from mmlab_pipeline.normalization import normalize_role, normalize_ranking
from mmlab_pipeline.parser import parse_and_validate_email


def paper(**values):
    return {**bundle()['reports'][0], 'id':'one', 'date':'2026-09-03','type':'Paper',
            'title':'A study of realistic video retrieval','venue':'MAPR 2026','ranking':'C-Unranked',
            'status':'Accepted','role':'Co-author','authors':'Huy Che','index':'Scopus',
            'memberIds':[7], 'missingFields':[],'invalidFields':[],'issues':[],**values}


@pytest.mark.parametrize('value,expected', [('Co-authors','Co-author'),('co authoor','Co-author'),
    ('first autthor','First author'),('đồng tác giả','Co-author'),('not co author','not co author'),
    ('author','author'),('corresponding author','corresponding author'),('XYZ','XYZ')])
def test_conservative_role_normalization(value,expected):
    assert normalize_role(value)==expected


def test_cleanup_preserves_evidence_and_source():
    original=paper(role='Co-authors',ranking='Unranked',status='Chưa rõ',subject='Fwd: Decision available',missingFields=['Title of Work'])
    before=copy.deepcopy(original);r=clean_report(original)
    assert original==before
    assert r['role']=='Co-author' and r['ranking']=='Unranked'
    assert r['memberIds']==before['memberIds'] and r['status']=='Chưa rõ'
    assert 'Title of Work' in r['missingFields'] and not r['isValid']
    assert normalize_ranking('A-ranked')=='A'
    assert clean_report(paper(status='Chưa rõ',subject='Paper is not accepted'))['status']=='Chưa rõ'
    e={'field':'Role','value':'First author','reason':'Trường xuất hiện nhiều lần; không tự chọn hoặc gộp giá trị'}
    r=clean_report(paper(invalidFields=[e]));assert not r['isValid'] and e in r['invalidFields']


def test_metadata_and_stale_validation():
    body='Type of Report: Paper\nTitle of Work: A study of video\nAll authors: Huy Che\nVenue: MAPR 2026\nRole: Co-authors\nIndex: Scopus\nRanking: C-unrank\nSubject: Decision'
    r=parse_and_validate_email('Fwd: Decision',body,'student@gmail.com')
    assert r['is_valid'] and r['data']['role']=='Co-author'
    assert parse_and_validate_email('Fwd: NCS',body.replace('Type of Report: Paper','Type of Report: NCS'),'huycq@uit.edu.vn')['is_valid']
    old=clean_report(paper(invalidFields=[{'field':'Subject','value':'Decision','reason':'Extra inputs are not permitted'}]))
    assert old['isValid']
    monthly=clean_report(paper(type='Báo cáo tháng',monthlyTasks={'done':None,'planned':None},invalidFields=[{'field':'Subject','value':'Whatever','reason':'Subject phải đúng mẫu'}]))
    assert monthly['isValid'] and monthly['reportPeriod']=='2026-09'
    assert monthly['monthlyTasks']=={'done':None,'planned':None}


def test_compose_deduplicates_and_uses_receipt_period():
    one=paper(taskGroup='strategic')
    two=paper(id='two',venue='MAPR2026',memberIds=[1],taskGroup='strategic')
    unknown=paper(id='unknown',title='A different paper title',status='Chưa rõ')
    later=paper(id='later',date='2026-10-01')
    deleted=paper(id='deleted',title='Deleted study',deletedAt='2026-09-04')
    external=paper(id='external',title='External research paper',memberIds=[])
    monthly=paper(id='monthly',type='Báo cáo tháng',taskGroup='routine',monthlyTasks={
        'done':[{'date':'','content':'Hoàn tất bộ dữ liệu'}], 'planned':['Sẽ tổ chức seminar']})
    r=compose([one,two,unknown,later,deleted,external,monthly],'2026-09')
    assert r['counts']['acceptedConference']==1 and r['counts']['unknownPaperStatus']==1
    assert r['sections']['a1'].count(one['title'])==1
    assert 'Nguyễn Vinh Tiệp' in r['sections']['a1'] and 'Chế Quang Huy' in r['sections']['a1']
<<<<<<< HEAD
    assert 'Được chấp nhận' not in r['sections']['a1']
    assert r['sections']['a1']==r['sections']['a2'] and r['sections']['b1']==r['sections']['b2']
=======
    assert 'Được chấp nhận 1 bài báo hội nghị.' in r['sections']['a1']
>>>>>>> 7d8b3d32e37d57bbb3f356c1bd45e50e3337c878
    assert 'Hoàn tất bộ dữ liệu' in r['sections']['a2'] and 'Sẽ tổ chức seminar' in r['sections']['b2']
    assert r['planPeriod']=='2026-10' and r['records']==5


def doc_text(content):return '\n'.join(p.text for p in Document(io.BytesIO(content)).paragraphs)


def test_docx_variants_and_blank_dates():
<<<<<<< HEAD
    r=compose([],'2026-12');r['sections'].update(a1='KPI bài báo: hoàn thành 02 bài báo Scopus',b1='KPI bài báo: nộp 03 bài tạp chí Q1')
    discussion=doc_text(docx_bytes(r,'discussion'));school=doc_text(docx_bytes(r,'school'))
    assert 'hoàn thành 02 bài báo Scopus' in discussion and 'Scopus: 2 bài.' in school
    assert 'Phần A.1' in school and 'Phần A.2' in school and 'nộp 03 bài tạp chí Q1' in school and 'THÁNG 01/2027' in school
=======
    r=compose([],'2026-12');r['sections'].update(a1='ĐÃ-HOÀN-TẤT',b1='SẼ-THỰC-HIỆN')
    discussion=doc_text(docx_bytes(r,'discussion'));school=doc_text(docx_bytes(r,'school'))
    assert 'ĐÃ-HOÀN-TẤT' in discussion and 'ĐÃ-HOÀN-TẤT' in school
    assert 'Phần A.1' in school and 'Phần A.2' in school and 'SẼ-THỰC-HIỆN' in school and 'THÁNG 01/2027' in school
>>>>>>> 7d8b3d32e37d57bbb3f356c1bd45e50e3337c878
    template=doc_text(blank_template('school'))
    assert 'THÁNG …/……' in template and 'Ngày … tháng … năm ……' in template


@pytest.mark.parametrize('year,month,day', [(2026,9,28),(2026,8,31),(2026,2,23),(2024,2,26),(2026,12,28),(2027,1,25)])
def test_last_monday(year,month,day):
    due=service.due_date(year,month)
    assert due.day==day and due.weekday()==0 and due.hour==9 and due.utcoffset().total_seconds()==7*3600


@pytest.fixture
def scheduled(client,monkeypatch):
    monkeypatch.setenv('MONTHLY_REPORTS_ENABLED','true');monkeypatch.setenv('MAIL_MODE','preview')
    with store.mutation() as c:
        for table in (service.deliveries,service.batches,service.drafts):c.execute(table.delete())
<<<<<<< HEAD
        c.execute(store.workflow_settings.delete().where(store.workflow_settings.c.key.like('school_kpi:%')))
=======
>>>>>>> 7d8b3d32e37d57bbb3f356c1bd45e50e3337c878
        c.execute(update(store.workflow_settings).where(store.workflow_settings.c.key=='monthly_started').values(value='2026-09-01T00:00:00+07:00'))
    yield client
    with store.mutation() as c:
        for table in (service.deliveries,service.batches,service.drafts):c.execute(table.delete())
<<<<<<< HEAD
        c.execute(store.workflow_settings.delete().where(store.workflow_settings.c.key.like('school_kpi:%')))
=======
>>>>>>> 7d8b3d32e37d57bbb3f356c1bd45e50e3337c878


def test_schedule_timezone_restart_catchup_and_no_duplicates(scheduled,monkeypatch):
    assert not service.enqueue_due(datetime(2026,9,28,1,59,tzinfo=timezone.utc))
    assert service.enqueue_due(datetime(2026,9,28,2,0,tzinfo=timezone.utc))
    assert not service.enqueue_due(datetime(2026,10,2,9,tzinfo=VN))
    with store.engine.connect() as c:rows=c.execute(select(service.deliveries)).mappings().all()
    assert len(rows)==12 and {r['recipient'] for r in rows}=={m['email'] for m in service.DIRECTORY}
    calls=[];service.dispatch(sender=lambda row:calls.append(row) or ('sent',None));assert calls==[]
    monkeypatch.setenv('MAIL_MODE','smtp')
    service.dispatch(sender=lambda row:calls.append(row) or ('sent',None))
    service.dispatch(sender=lambda row:calls.append(row) or ('sent',None))
    assert len(calls)==12
    for row in calls:
        assert len(row['payload']['attachments'])==2
        for a in row['payload']['attachments']:Document(io.BytesIO(base64.b64decode(a['content'])))


def test_catchup_activation_and_uncertain_delivery(scheduled,monkeypatch):
    assert not service.enqueue_due(datetime(2026,9,10,9,tzinfo=VN))
    assert service.enqueue_due(datetime(2026,10,2,9,tzinfo=VN))
    monkeypatch.setenv('MAIL_MODE','smtp');calls=[]
    service.dispatch(sender=lambda row:calls.append(row) or ('unknown','timeout'))
    service.dispatch(sender=lambda row:calls.append(row) or ('sent',None));assert len(calls)==12
    with store.mutation() as c:c.execute(update(service.deliveries).values(status='sending'))
    service.recover_interrupted()
    with store.engine.connect() as c:assert set(c.execute(select(service.deliveries.c.status)).scalars())=={'unknown'}


def test_disabled_schedule(scheduled,monkeypatch):
    monkeypatch.setenv('MONTHLY_REPORTS_ENABLED','false')
    assert not service.enqueue_due(datetime(2026,9,28,9,tzinfo=VN))


def test_admin_monthly_edit_conflicts_download_and_saved_snapshot(scheduled):
    c=scheduled
    assert c.get('/api/monthly/2026-09').status_code==401
    assert c.get('/api/monthly/2026-09/school.docx').status_code==401
    login(c);assert c.get('/api/monthly/2026-99').status_code==422
    r=c.get('/api/monthly/2026-09').json()
<<<<<<< HEAD
    body={k:r[k] for k in ('revision','strategyLabel','signatory','sections')};body['sections']['b1']='01 NCS sẽ báo cáo chuyên đề 3'
=======
    body={k:r[k] for k in ('revision','strategyLabel','signatory','sections')};body['sections']['b1']='Sẽ báo cáo seminar đã thống nhất'
>>>>>>> 7d8b3d32e37d57bbb3f356c1bd45e50e3337c878
    assert c.put('/api/monthly/2026-09',json=body).status_code==403
    assert c.put('/api/monthly/2026-09',json=body,headers={'Origin':'http://testserver'}).status_code==200
    assert c.put('/api/monthly/2026-09',json=body,headers={'Origin':'http://testserver'}).status_code==409
    assert c.get('/api/monthly/2026-09').json()['sections']['b1']==body['sections']['b1']
    assert c.get('/api/monthly/2026-09?fresh=true').json()['sections']['b1']==''
    response=c.get('/api/monthly/2026-09/school.docx');assert response.status_code==200 and response.content[:2]==b'PK'
    assert service.enqueue_due(datetime(2026,9,28,9,tzinfo=VN))
    with store.engine.connect() as db:p=db.execute(select(service.batches.c.payload)).scalar_one()
    assert body['sections']['b1'] in doc_text(base64.b64decode(p['attachments'][1]['content']))


def test_smtp_two_attachments(scheduled,monkeypatch):
    service.enqueue_due(datetime(2026,9,28,9,tzinfo=VN))
    with store.engine.connect() as c:
        row=dict(c.execute(select(service.deliveries)).mappings().first());row['payload']=c.execute(select(service.batches.c.payload)).scalar_one()
    monkeypatch.setenv('MMLAB_APP_PASSWORD','synthetic-test-only');captured={}
    class SMTP:
        def __init__(self,*args,**kwargs):pass
        def login(self,*args):pass
        def send_message(self,message,from_addr,to_addrs):captured.update(message=message,to=to_addrs)
        def quit(self):pass
    with patch.object(mailer.smtplib,'SMTP_SSL',SMTP):assert mailer.deliver(row)==('sent',None)
    assert captured['to']==[row['recipient']] and len(list(captured['message'].iter_attachments()))==2
