# -*- coding: utf-8 -*-
"""MMLab monthly form: two inputs A and B, school output A.1=A.2 and B.1=B.2.

Apply AFTER update_mmlab_auto_receipt.py from the previous message.
Place beside run_local.py and stop the app before updating:
    python update_mmlab_form_ab.py
    npm run build --prefix frontend
    python run_local.py

Use --check to preview compatibility without writing any files.
Backups are under _backups/paired-ab-*. No network calls or email sending.
No database, .env, credentials or source email records are changed by this script.

The form shows only the two body fields. Publication/confirmation status labels
are not generated in the monthly report. Stored publication metadata is preserved.
School DOCX repeats A twice and B twice; discussion DOCX includes each once.
Older saved drafts merge all distinct A/B content on read without overwriting the
saved original. Clearing A also clears A.2 (likewise B). Source counts stay unique.
Legacy recommendations remain in saved payloads but aren't displayed/exported.
The existing automatic approval/receipt email workflow is unchanged.

This is plain, reviewable source code with an embedded unified diff; no ZIP needed.
"""
from pathlib import Path
from datetime import datetime
import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile

MANIFEST = {'backend/monthly_assembly.py': {'before': '66457df929251e63b29a3127c4d929f53db16d9710a60744235f8c41f9a808ae', 'after': '406cb9145a5686a36b0873f32313ac2af3654a9cf44501c4da00a10c2adf7b5e'}, 'backend/monthly_layout.py': {'before': None, 'after': 'b9e093ca2be201b4a61ea1524b403cd38bbdffebb90d54b23921833892ef7ae0'}, 'backend/monthly_report.py': {'before': 'e816f7b64bd655de486730ab8d35ecb5af3d026b7352b0952478cf3923a4837e', 'after': 'a506b9c6232148a553f958c42d07cd40719de70ccb1169b9618f0e2e32ed7275'}, 'backend/monthly_service.py': {'before': '11ae3b8091e21cb433318243267d0bfd0cf02443afb2f9623bdb372dbef2ffc6', 'after': '484f86c64eef8091bbaa22d6478050839c2954a297839a7a9bd0d3ebe847c5ce'}, 'backend/tests/test_auto_receipt.py': {'before': '16039c1cbeef045996d12213e9b73f2f8ac3772ed64445403767cb937010c378', 'after': '3178ddc6276272413760917bae19b3813456664cf1a1a1f5d9e3f882940c0141'}, 'backend/tests/test_monthly.py': {'before': '9abe3112ea6539692feebf19da7a6fd02a59a8a9fd0558cffad4650e74297383', 'after': '54323ce9c62666d6a3b7e2b0c3255914d2c038b337ce436c67b0ce33dde7edea'}, 'backend/tests/test_paired_monthly.py': {'before': None, 'after': '1e18d2a68d2e9ab48f2e22a12dbe9ef8e22733581b9a2acd1b2c04d04f99c1ed'}, 'backend/tests/test_task_assembly.py': {'before': '00062024784ab0f916666073b713370f5c621904e8271dc48889d4b126a88619', 'after': '76dd01ef0745f709e159fd069dbd2d9f12a1d51419174c6cd6dda28b8f813224'}, 'frontend/src/MonthlyReports.tsx': {'before': '8dfd84b8b1ccaf55ac33fd9ff87a4e051448efe7caa0e4ddd07bcfd19c9ad013', 'after': '4a087bf6cf6253903ebd72aef8455843afbeb279b5c71e9872ea885b2688b461'}}

PATCH = r'''--- a/backend/monthly_assembly.py
+++ b/backend/monthly_assembly.py
@@ -28,9 +28,8 @@
     records=[r for r in records if not r.get('deletedAt') and r.get('date','').startswith(period)]
     entries={key:[] for key in section_keys};review=[];seen=set();seen_tasks=set()
     counts={'acceptedConference':0,'acceptedJournal':0,'submittedConference':0,'submittedJournal':0,'unknownPaperStatus':0}
-    grouped_counts={g:{k:0 for k in counts} for g in ('1','2','Other')}
     for r in sorted(records,key=lambda x:(x['date'],x['id'])):
-        group={'strategic':'1','routine':'2'}.get(r.get('taskGroup'),'Other')
+        group='1'
         member_ids=set(r.get('memberIds',[]))
         names=', '.join(m['name'] for m in DIRECTORY if m['id'] in member_ids)
         if r['type']=='Báo cáo tháng':
@@ -54,31 +53,21 @@
             copies=[x for x in records if work_key(x)==key]
             member_ids={mid for x in copies for mid in x.get('memberIds',[])}
             names=', '.join(m['name'] for m in DIRECTORY if m['id'] in member_ids)
-            groups={x.get('taskGroup','unassigned') for x in copies}
-            group={'strategic':'1','routine':'2'}.get(next(iter(groups)),'Other') if len(groups)==1 else 'Other'
             accepted=any(fold(x.get('status','')) in ('accepted','accept','duoc chap nhan','da chap nhan') for x in copies)
             submitted=any(fold(x.get('status','')) in ('submitted','da nop') for x in copies)
-            action='Được chấp nhận' if accepted else 'Đã nộp' if submitted else 'Bài báo chưa rõ trạng thái'
             rank=r.get('ranking','')
             kind='Journal' if rank in ('Q1','Q2','Q3','Q4','Q-Unranked') else 'Conference' if rank in ('A*','A','B','C','C-Unranked') else ''
             bucket='accepted'+kind if accepted and kind else 'submitted'+kind if submitted and kind else None
             if bucket and member_ids:
-                counts[bucket]+=1;grouped_counts[group][bucket]+=1
+                counts[bucket]+=1
             elif not accepted and not submitted:counts['unknownPaperStatus']+=1
-            text=f"{action}: {r['title']} — {r.get('venue') or 'chưa có venue'}; {rank or 'chưa có ranking'}; {r.get('index') or 'chưa có index'}"
+            text=f"{r['title']} — {r.get('venue') or 'chưa có venue'}; {rank or 'chưa có ranking'}; {r.get('index') or 'chưa có index'}"
         else:text=f"{r['type']}: {r['title']}"
         if names:entries['a'+group].append((topic(text,r['type']),f'{text} ({names})'))
         else:review.append('Chưa đối chiếu được thành viên lab: '+text)
-    labels={'acceptedConference':'Được chấp nhận {n} bài báo hội nghị',
-            'acceptedJournal':'Được chấp nhận {n} bài báo tạp chí',
-            'submittedConference':'Đã nộp {n} bài báo hội nghị',
-            'submittedJournal':'Đã nộp {n} bài báo tạp chí'}
     sections={}
     for key,items in entries.items():
         lines=[]
-        if key.startswith('a'):
-            values=grouped_counts[key[1:]]
-            lines += [labels[k].format(n=values[k])+'.' for k in labels if values[k]]
         for heading in TOPICS:
             content=list(dict.fromkeys(text for category,text in items if category==heading))
             if content:lines+=[heading+':']+content
--- a/backend/monthly_layout.py
+++ b/backend/monthly_layout.py
@@ -0,0 +1,53 @@
+"""Two editable bodies; repeat them in the school's four required slots."""
+import re
+from monthly_assembly import TOPICS
+
+LAYOUT = 'paired-ab/1'
+LEGACY_PREFIX = re.compile(
+    r'^(?:Được chấp nhận|Đã nộp|Bài báo chưa rõ trạng thái|'
+    r'Ghi nhận bài báo, chưa có trạng thái xác định|Accepted|Chưa xác nhận)\s*:\s*', re.I)
+LEGACY_TOTAL = re.compile(r'^(?:Được chấp nhận|Đã nộp) \d+ bài báo (?:hội nghị|tạp chí)\.$')
+
+
+def merge_legacy(blocks):
+    """Keep unique legacy content, grouping repeated headings only once."""
+    groups = {'': []}
+    seen = set()
+    for block in blocks:
+        group = ''
+        for raw in block.splitlines():
+            line = raw.strip()
+            if not line or LEGACY_TOTAL.fullmatch(line):
+                continue
+            line = LEGACY_PREFIX.sub('', line)
+            if line.rstrip(':') in TOPICS:
+                group = line.rstrip(':')
+                groups.setdefault(group, [])
+                continue
+            if line not in seen:
+                seen.add(line)
+                groups.setdefault(group, []).append(line)
+    result = []
+    for group, lines in groups.items():
+        if not lines:
+            continue
+        if group:
+            result.append(group + ':')
+        result.extend(lines)
+    return '\n'.join(result)
+
+
+def paired_report(report):
+    """Non-mutating compatibility layer for saved drafts and old exports."""
+    result = {**report, 'sections': dict(report['sections']), 'layoutVersion': LAYOUT}
+    sections = result['sections']
+    for phase in ('a', 'b'):
+        if report.get('layoutVersion') == LAYOUT:
+            body = sections.get(phase + '1', '')  # Empty means intentionally cleared.
+        else:
+            body = merge_legacy([sections.get(phase + suffix, '') for suffix in ('1', '2', 'Other')])
+        sections[phase + '1'] = sections[phase + '2'] = body
+        sections[phase + 'Other'] = ''
+    # Keep legacy recommendations in stored payloads, but don't render a third part.
+    sections.setdefault('recommendations', '')
+    return result
--- a/backend/monthly_report.py
+++ b/backend/monthly_report.py
@@ -8,6 +8,7 @@
 from docx.oxml.ns import qn
 from mmlab_pipeline.members import fold
 from monthly_assembly import assemble, TOPICS
+from monthly_layout import paired_report, LAYOUT
 
 VN=timezone(timedelta(hours=7))
 SECTIONS=('a1','a2','aOther','b1','b2','bOther','recommendations')
@@ -37,20 +38,13 @@
 
 def compose(records, period):
     validate_period(period)
-    return assemble(records,period,SECTIONS,DEFAULT_STRATEGY,work_key)
+    return paired_report({**assemble(records,period,SECTIONS,DEFAULT_STRATEGY,work_key),'layoutVersion':LAYOUT})
 
 
 def docx_bytes(report, variant):
     if variant not in ('discussion','school'):raise ValueError('Unknown report variant')
-    # Work on a copy: rendering must not mutate the admin's saved draft.
-    report={**report,'sections':dict(report['sections'])}
+    report=paired_report(report)
     if report['strategyLabel']=='KHCL Trường giai đoạn 2021-2025':report['strategyLabel']=DEFAULT_STRATEGY
-    if variant=='school':
-        # Only explicitly tagged strategic tasks enter A.1/B.1.
-        # Other tasks are regular work; preserve their source text.
-        for phase in ('a','b'):
-            report['sections'][phase+'2']='\n'.join(filter(None,[report['sections'].get(phase+'2','').strip(),report['sections'].get(phase+'Other','').strip()]))
-            report['sections'][phase+'Other']=''
     doc=Document(); sec=doc.sections[0]
     sec.page_width=Cm(21);sec.page_height=Cm(29.7)
     sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2)
@@ -81,7 +75,6 @@
     if not template:
         paragraph(f"Kết quả tháng {report['period'][5:]}/{report['period'][:4]} và kế hoạch tháng {m:02d}/{y}",center=True)
     def section(key,title):
-        if variant!='school' and not template and not report['sections'].get(key,'').strip():return
         if template and key=='b1' and variant=='discussion':doc.add_page_break()
         doc.add_heading(title,level=1)
         if key in ('a1','b1','b2'):
@@ -96,18 +89,14 @@
                     p.paragraph_format.left_indent=Cm(.35)
                     p.paragraph_format.first_line_indent=Cm(-.35)
                     p.runs[0].text='• '+line
-    if variant in ('discussion','school'):
+    if variant=='school':
         section('a1',f"Phần A.1: Tình hình thực hiện nhiệm vụ phục vụ {report['strategyLabel']} trong tháng trước")
         section('a2','Phần A.2: Tình hình thực hiện nhiệm vụ thường xuyên và đột xuất trong tháng trước')
-        if report['sections'].get('aOther'):section('aOther', 'Phần A: Kết quả công tác trong tháng trước' if not report['sections'].get('a1') and not report['sections'].get('a2') else 'Nội dung thực hiện bổ sung')
-    section('b1',f"Phần B.1: Kế hoạch công tác các nhiệm vụ phục vụ {report['strategyLabel']} trong tháng này")
-    section('b2','Phần B.2: Kế hoạch công tác các nhiệm vụ thường xuyên và đột xuất trong tháng này')
-    if report['sections'].get('bOther'):section('bOther','Phần B: Kế hoạch công tác trong tháng này' if not report['sections'].get('b1') and not report['sections'].get('b2') else 'Nhiệm vụ dự kiến bổ sung')
-    if variant=='discussion' and report.get('reviewNotes'):
-        doc.add_heading('Nội dung cần bổ sung để hoàn thiện báo cáo',level=1)
-        for note in report['reviewNotes']:paragraph(note)
-    if variant=='school' or template or report['sections'].get('recommendations','').strip():
-        section('recommendations','Phần 3: Các kiến nghị')
+        section('b1',f"Phần B.1: Kế hoạch công tác các nhiệm vụ phục vụ {report['strategyLabel']} trong tháng này")
+        section('b2','Phần B.2: Kế hoạch công tác các nhiệm vụ thường xuyên và đột xuất trong tháng này')
+    else:
+        section('a1','Phần A: Công việc đã thực hiện')
+        section('b1','Phần B: Kế hoạch công việc')
     if doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
     p=paragraph('Trưởng đơn vị\n(Ký và ghi rõ họ tên)\n\n'+report['signatory'],True)
     p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
@@ -119,19 +108,7 @@
 def blank_template(variant):
     report=compose([], '2026-09')
     report['isTemplate']=True
-    done=['Được chấp nhận … bài báo hội nghị rank A*/A/B/C/Scopus',
-          'Được chấp nhận … bài báo tạp chí ISI/Scopus Q1/Q2/Q3/Q4',
-          'Đã nộp … bài báo hội nghị rank A*/A/B/C/Scopus',
-          'Đã nộp … bài báo tạp chí ISI/Scopus Q1/Q2/Q3/Q4',
-          'Hoàn tất đăng ký … đề tài NCKH cấp cơ sở/cấp ĐHQG/NAFOSTED',
-          'Hoàn tất nghiệm thu … đề tài NCKH cấp cơ sở/cấp ĐHQG/NAFOSTED',
-          '… NCS đã hoàn thành nhập học/báo cáo CĐ 1/CĐ 2/CĐ 3/TLTQ/Seminar/ĐVCM/Cấp Trường',
-          'Đã thực hiện báo cáo … seminar học thuật tại PTN', 'Đã đạt thành tích giải thưởng …']
-    plans=['Sẽ nộp … bài báo hội nghị rank A*/A/B/C/Scopus',
-           'Sẽ nộp … bài báo tạp chí ISI/Scopus Q1/Q2/Q3/Q4',
-           'Sẽ đăng ký … đề tài NCKH cấp cơ sở/cấp ĐHQG/NAFOSTED',
-           'Sẽ nghiệm thu … đề tài NCKH cấp cơ sở/cấp ĐHQG/NAFOSTED',
-           '… NCS sẽ nhập học/báo cáo CĐ 1/CĐ 2/CĐ 3/TLTQ/Seminar/ĐVCM/Cấp Trường',
-           'Sẽ thực hiện … seminar học thuật tại PTN', 'Dự kiến tham gia giải thưởng/cuộc thi …']
+    done=['Tên công trình hoặc nội dung công việc; kết quả; người thực hiện.']
+    plans=['Nội dung công việc dự kiến; thời gian; người thực hiện.']
     report['sections'].update(a1='\n'.join(done),a2='\n'.join(done),b1='\n'.join(plans),b2='\n'.join(plans),recommendations='…')
     return docx_bytes(report,variant)
--- a/backend/monthly_service.py
+++ b/backend/monthly_service.py
@@ -12,6 +12,7 @@
 import mailer
 from monthly_report import VN, SECTIONS, compose, docx_bytes, validate_period, DEFAULT_STRATEGY
 from mmlab_pipeline.members import DIRECTORY
+from monthly_layout import paired_report, LAYOUT
 
 drafts=Table('monthly_drafts',store.meta,Column('period',String,primary_key=True),Column('payload',JSON,nullable=False),Column('revision',Integer,nullable=False))
 batches=Table('monthly_batches',store.meta,Column('period',String,primary_key=True),Column('payload',JSON,nullable=False),Column('created_at',String,nullable=False))
@@ -56,19 +57,21 @@
     if saved and not fresh:
         report={**saved['payload'],'revision':saved['revision'],'customized':True}
         if report['strategyLabel']=='KHCL Trường giai đoạn 2021-2025':report['strategyLabel']=DEFAULT_STRATEGY
-        return report
+        return paired_report(report)
     records=conn.execute(select(store.reports.c.payload).where(store.reports.c.mailbox==store.mailbox())).scalars().all()
     return {**compose(records,period),'revision':saved['revision'] if saved else 0,'customized':False}
 
 
 def save_draft(period,data):
     validate_period(period)
-    if set(data.sections)!=set(SECTIONS) or any(len(v)>100000 for v in data.sections.values()):
+    if set(data.sections) not in ({'a1','b1'},set(SECTIONS)) or any(len(v)>100000 for v in data.sections.values()):
         raise HTTPException(422,'Nội dung các phần chưa đúng hoặc quá dài.')
     with store.mutation() as conn:
         report=read_draft(period,conn)
         if data.revision!=report['revision']:raise HTTPException(409,'Báo cáo tháng vừa thay đổi. Hãy tải lại.')
-        report.update(strategyLabel=data.strategyLabel,signatory=data.signatory,sections=data.sections,revision=data.revision+1,customized=True)
+        sections={**report['sections'],'a1':data.sections['a1'],'b1':data.sections['b1']}
+        report.update(strategyLabel=data.strategyLabel,signatory=data.signatory,sections=sections,layoutVersion=LAYOUT,revision=data.revision+1,customized=True)
+        report=paired_report(report)
         stmt=store.insert(drafts).values(period=period,payload=report,revision=report['revision'])
         conn.execute(stmt.on_conflict_do_update(index_elements=['period'],set_={'payload':stmt.excluded.payload,'revision':stmt.excluded.revision}))
     return report
@@ -87,7 +90,7 @@
         attachments=[{'name':f'mmlab-{report["planPeriod"]}-{variant}.docx','content':base64.b64encode(docx_bytes(report,variant)).decode(),
                       'subtype':'vnd.openxmlformats-officedocument.wordprocessingml.document'} for variant in ('discussion','school')]
         payload={'subject':f'[MMLab] Báo cáo thảo luận và kế hoạch tháng {report["planPeriod"]}',
-                 'body':f'Kính gửi các thành viên MMLab,\n\nĐính kèm hai bản báo cáo:\n1. Bản thảo luận: công việc đã thực hiện tháng {period} và kế hoạch tháng {report["planPeriod"]}.\n2. Bản báo cáo nộp trường: A.1, A.2, B.1, B.2 và kiến nghị; đơn vị mã 6, chiến lược 2021–2030.\n\nVui lòng trao đổi và bổ sung các nội dung chưa phân nhóm hoặc còn thiếu.\n\nTrân trọng,\nMMLab — UIT',
+                 'body':f'Kính gửi các thành viên MMLab,\n\nĐính kèm hai bản báo cáo:\n1. Bản thảo luận: công việc đã thực hiện tháng {period} và kế hoạch tháng {report["planPeriod"]}.\n2. Bản báo cáo nộp trường: A.1 = A.2, B.1 = B.2; đơn vị mã 6, chiến lược 2021–2030.\n\nVui lòng trao đổi về kết quả và kế hoạch công việc.\n\nTrân trọng,\nMMLab — UIT',
                  'attachments':attachments}
         conn.execute(store.insert(batches).values(period=period,payload=payload,created_at=now.isoformat()))
         assert len(DIRECTORY)==12 and len({m['email'] for m in DIRECTORY})==12
--- a/backend/tests/test_auto_receipt.py
+++ b/backend/tests/test_auto_receipt.py
@@ -83,13 +83,14 @@
 
 def test_school_form_sections_unit_strategy_month_and_no_invented_targets():
     r=compose([paper(taskGroup='unassigned',date='2026-07-15')],'2026-07')
-    r['sections'].update(bOther='Chuẩn bị hồ sơ nghiên cứu sinh',recommendations='Đề nghị hỗ trợ máy tính')
+    r['sections'].update(b1='Chuẩn bị hồ sơ nghiên cứu sinh',recommendations='Đề nghị hỗ trợ máy tính')
     original=copy.deepcopy(r)
     text=doc_text(docx_bytes(r,'school'))
-    for expected in ('THÁNG 08/2026','Mã đơn vị: 6','Phần A.1','Phần A.2','Phần B.1','Phần B.2','Phần 3','2021-2030','KH2026','Đề nghị hỗ trợ máy tính'):
+    for expected in ('THÁNG 08/2026','Mã đơn vị: 6','Phần A.1','Phần A.2','Phần B.1','Phần B.2','2021-2030','KH2026'):
         assert expected in text
-    assert text.count('A study of realistic video retrieval')==1
+    assert text.count('A study of realistic video retrieval')==2
     a1=text.split('Phần A.1')[1].split('Phần A.2')[0]
-    assert 'A study of realistic video retrieval' not in a1
+    assert 'A study of realistic video retrieval' in a1
+    assert 'Phần 3' not in text and 'Đề nghị hỗ trợ máy tính' not in text
     assert '46 / 44' not in text and '2021-2025' not in text
     assert r==original  # School output must not modify saved discussion data.
--- a/backend/tests/test_monthly.py
+++ b/backend/tests/test_monthly.py
@@ -68,7 +68,8 @@
     assert r['counts']['acceptedConference']==1 and r['counts']['unknownPaperStatus']==1
     assert r['sections']['a1'].count(one['title'])==1
     assert 'Nguyễn Vinh Tiệp' in r['sections']['a1'] and 'Chế Quang Huy' in r['sections']['a1']
-    assert 'Được chấp nhận 1 bài báo hội nghị.' in r['sections']['a1']
+    assert 'Được chấp nhận' not in r['sections']['a1']
+    assert r['sections']['a1']==r['sections']['a2'] and r['sections']['b1']==r['sections']['b2']
     assert 'Hoàn tất bộ dữ liệu' in r['sections']['a2'] and 'Sẽ tổ chức seminar' in r['sections']['b2']
     assert r['planPeriod']=='2026-10' and r['records']==5
 
--- a/backend/tests/test_paired_monthly.py
+++ b/backend/tests/test_paired_monthly.py
@@ -0,0 +1,80 @@
+import copy
+import io
+from docx import Document
+from sqlalchemy import select
+from monthly_layout import paired_report, LAYOUT
+from monthly_report import compose, docx_bytes
+import monthly_service as service
+from test_monthly import paper, scheduled, client, doc_text
+from test_api import login
+
+
+def bodies(blob):
+    result = {}; key = None
+    for p in Document(io.BytesIO(blob)).paragraphs:
+        if p.style.name == 'Heading 1':
+            key = p.text.split(':')[0];result[key] = []
+        elif key and p.text.startswith('• '):
+            result[key].append(p.text)
+    return result
+
+
+def test_mirrored_output_without_doubling_source_counts_or_status_labels():
+    rows=[paper(status='Accepted'),paper(id='two',title='Another scientific article',status='Chưa rõ')]
+    r=compose(rows,'2026-09')
+    assert r['works']==2 and r['counts']['acceptedConference']==1
+    assert r['sections']['a1']==r['sections']['a2']
+    assert r['sections']['a1'].count('A study of realistic video retrieval')==1
+    school=bodies(docx_bytes(r,'school'))
+    assert school['Phần A.1']==school['Phần A.2']
+    assert school['Phần B.1']==school['Phần B.2']
+    for variant in ('school','discussion'):
+        text=doc_text(docx_bytes(r,variant))
+        for label in ('Accepted','Được chấp nhận','chưa rõ trạng thái','Chưa xác nhận','Phần 3'):
+            assert label not in text
+    assert len(bodies(docx_bytes(r,'discussion')))==2
+    assert rows[0]['status']=='Accepted' and rows[1]['status']=='Chưa rõ'
+
+
+def test_legacy_draft_merge_preserves_unique_tasks_and_no_duplicate_headings():
+    r=compose([],'2026-09');r.pop('layoutVersion')
+    r['sections'].update(a1='Công bố khoa học:\nĐược chấp nhận 1 bài báo hội nghị.\nĐược chấp nhận: Paper Alpha (Thìn)',
+        a2='Công bố khoa học:\nĐược chấp nhận: Paper Alpha (Thìn)\nAccepted: Paper Beta (Huy)',
+        aOther='Công bố khoa học:\nBài báo chưa rõ trạng thái: Paper Gamma (Tiệp)',
+        b1='Dự định một',b2='Dự định hai',bOther='Dự định ba',recommendations='Kiến nghị cũ')
+    original=copy.deepcopy(r)
+    converted=paired_report(r)
+    a=converted['sections']['a1']
+    assert converted['sections']['a2']==a and a.count('Công bố khoa học:')==1
+    assert all(a.count(s)==1 for s in ('Paper Alpha','Paper Beta','Paper Gamma'))
+    assert converted['sections']['b1']==converted['sections']['b2']=='Dự định một\nDự định hai\nDự định ba'
+    assert converted['sections']['recommendations']=='Kiến nghị cũ'
+    assert paired_report(converted)==converted and r==original
+    converted['sections']['a1']=''
+    assert paired_report(converted)['sections']['a2']==''
+
+
+def test_two_field_api_save_download_and_scheduled_attachments(scheduled):
+    c=scheduled;login(c);r=c.get('/api/monthly/2026-09').json()
+    payload={k:r[k] for k in ('revision','strategyLabel','signatory')}
+    payload['sections']={'a1':'Công việc tháng trước','b1':'Kế hoạch tháng này'}
+    response=c.put('/api/monthly/2026-09',json=payload,headers={'Origin':'http://testserver'})
+    assert response.status_code==200,response.text
+    sections=response.json()['sections']
+    assert sections['a1']==sections['a2']==payload['sections']['a1']
+    assert sections['b1']==sections['b2']==payload['sections']['b1']
+    assert c.put('/api/monthly/2026-09',json=payload,headers={'Origin':'http://testserver'}).status_code==409
+    text=doc_text(c.get('/api/monthly/2026-09/school.docx').content)
+    assert text.count('Công việc tháng trước')==2 and text.count('Kế hoạch tháng này')==2
+    from datetime import datetime
+    from monthly_report import VN
+    import base64, store
+    assert service.enqueue_due(datetime(2026,9,28,9,tzinfo=VN))
+    with store.engine.connect() as conn:
+        outgoing=conn.execute(select(service.batches.c.payload)).scalar_one()
+    text=doc_text(base64.b64decode(outgoing['attachments'][1]['content']))
+    assert text.count('Công việc tháng trước')==2
+    # Clearing A must also clear its mirrored A.2, not restore an earlier copy.
+    payload.update(revision=response.json()['revision'],sections={'a1':'','b1':'Kế hoạch mới'})
+    cleared=c.put('/api/monthly/2026-09',json=payload,headers={'Origin':'http://testserver'}).json()
+    assert cleared['sections']['a1']==cleared['sections']['a2']==''
--- a/backend/tests/test_task_assembly.py
+++ b/backend/tests/test_task_assembly.py
@@ -31,7 +31,7 @@
     a=paper(id='a',type='Báo cáo tháng',memberIds=[7],monthlyTasks={'done':[common]})
     b=paper(id='b',type='Báo cáo tháng',memberIds=[3],monthlyTasks={'done':[common]})
     result=compose([a,b],'2026-09')
-    assert result['sections']['aOther'].count('Thực hiện đề tài nghiên cứu')==2
+    assert result['sections']['a1'].count('Thực hiện đề tài nghiên cứu')==2
 
 
 def test_bullet_and_action_boundaries_not_joined():
@@ -40,7 +40,7 @@
     assert repair_tasks(tasks,join_wrapped=False)==tasks
     r=paper(type='Báo cáo tháng',monthlyTasksVersion=2,monthlyTasks={'done':tasks})
     result=compose([r],'2026-09')
-    assert result['sections']['aOther'].count('(Chế Quang Huy)')==3
+    assert result['sections']['a1'].count('(Chế Quang Huy)')==3
 
 
 def test_monthly_parser_does_not_drop_seminar_or_continuation_labels():
@@ -59,7 +59,7 @@
     result=compose([unknown,known],'2026-09')
     assert len(result['reviewNotes'])==1
     texts={v:'\n'.join(p.text for p in Document(io.BytesIO(docx_bytes(result,v))).paragraphs) for v in ('discussion','school')}
-    assert 'unknown@example.com' in texts['discussion'] and 'unknown@example.com' not in texts['school']
+    assert 'unknown@example.com' not in texts['discussion'] and 'unknown@example.com' not in texts['school']
     assert 'Phần A.1' in texts['school'] and 'Robotics' in texts['school']
     assert 'Chưa có nội dung báo cáo' not in texts['school']
     assert 'chưa phân nhóm' not in texts['school']
--- a/frontend/src/MonthlyReports.tsx
+++ b/frontend/src/MonthlyReports.tsx
@@ -2,23 +2,21 @@
 import {readApiResponse} from '@/lib/api';
 import {Button} from '@/components/ui/button';
 type Draft={period:string;planPeriod:string;revision:number;strategyLabel:string;signatory:string;sections:Record<string,string>;records:number;counts:Record<string,number>;reviewNotes?:string[];assemblyVersion?:number;customized?:boolean};
-const labels:Record<string,string>={a1:'A.1 · Đã — nhiệm vụ KHCL',a2:'A.2 · Đã — nhiệm vụ thường xuyên',aOther:'Đã — nội dung tổng hợp chung',b1:'B.1 · Sẽ — nhiệm vụ KHCL',b2:'B.2 · Sẽ — nhiệm vụ thường xuyên',bOther:'Sẽ — nội dung tổng hợp chung',recommendations:'Các kiến nghị'};
+const labels:Record<string,string>={a1:'Phần A · Công việc đã thực hiện (A.1 = A.2)',b1:'Phần B · Kế hoạch công việc (B.1 = B.2)'};
 export default function MonthlyReports(){
  const [period,setPeriod]=useState(new Intl.DateTimeFormat('sv-SE',{timeZone:'Asia/Ho_Chi_Minh',year:'numeric',month:'2-digit'}).format(new Date()));
  const [draft,setDraft]=useState<Draft|null>(null),[error,setError]=useState(''),[notice,setNotice]=useState(''),[busy,setBusy]=useState(false),[dirty,setDirty]=useState(false);
  const [next,setNext]=useState('');
  useEffect(()=>{let alive=true;setDraft(null);setError('');setDirty(false);fetch('/api/monthly/'+period).then(readApiResponse).then(d=>{if(alive)setDraft(d)}).catch(e=>{if(alive)setError(e.message)});return()=>{alive=false}},[period]);
  useEffect(()=>{fetch('/api/monthly/schedule').then(readApiResponse).then(d=>setNext(d.enabled?new Date(d.nextRun).toLocaleString('vi-VN',{timeZone:'Asia/Ho_Chi_Minh'}):'Đã tắt')).catch(()=>{});},[]);
- async function save(){if(!draft)return;setBusy(true);setError('');try{const d=await fetch('/api/monthly/'+period,{method:'PUT',headers:{'Content-Type':'application/json'},body:JSON.stringify({revision:draft.revision,strategyLabel:draft.strategyLabel,signatory:draft.signatory,sections:draft.sections})}).then(readApiResponse);setDraft(d);setDirty(false);setNotice('Đã lưu hai bản báo cáo.')}catch(e){setError(e instanceof Error?e.message:'Không lưu được.')}finally{setBusy(false)}}
+ async function save(){if(!draft)return;setBusy(true);setError('');try{const d=await fetch('/api/monthly/'+period,{method:'PUT',headers:{'Content-Type':'application/json'},body:JSON.stringify({revision:draft.revision,strategyLabel:draft.strategyLabel,signatory:draft.signatory,sections:{a1:draft.sections.a1,b1:draft.sections.b1}})}).then(readApiResponse);setDraft(d);setDirty(false);setNotice('Đã lưu hai bản báo cáo.')}catch(e){setError(e instanceof Error?e.message:'Không lưu được.')}finally{setBusy(false)}}
  async function refresh(){if(dirty&&!window.confirm('Bỏ thay đổi chưa lưu và tổng hợp lại từ Gmail?'))return;try{setDraft(await fetch('/api/monthly/'+period+'?fresh=true').then(readApiResponse));setDirty(true);setNotice('Đã tổng hợp lại. Bấm Lưu để dùng bản này khi tải hoặc gửi.')}catch(e){setError(e instanceof Error?e.message:'Không tải được.')}}
  return <section className="monthly-workspace"><div className="monthly-toolbar"><label>Tháng tổng hợp <input aria-label="Tháng tổng hợp" type="month" value={period} onChange={e=>{if(e.target.value&&(!dirty||window.confirm('Bỏ thay đổi chưa lưu và đổi tháng?')))setPeriod(e.target.value)}}/></label><Button variant="outline" onClick={()=>void refresh()}>Tổng hợp lại</Button><Button onClick={()=>void save()} disabled={!draft||busy||!dirty}>{busy?'Đang lưu…':'Lưu báo cáo'}</Button></div>
  <p className="monthly-schedule">Hai bản được gửi đến 12 thành viên lúc 09:00 thứ Hai cuối cùng mỗi tháng. Lịch tiếp theo: {next||'Đang tải…'}.</p>
  {error&&<p className="login-error" role="alert">{error}</p>}{notice&&<p role="status">{notice}</p>}
- {draft&&<>{draft.customized&&(!draft.assemblyVersion||draft.assemblyVersion<2)&&<p className="editor-hint">Đây là bản đã lưu bằng cách tổng hợp cũ. Bấm Tổng hợp lại để nối dòng và gộp mục trùng; kiểm tra nội dung trước khi lưu.</p>}<div className="monthly-downloads"><div><h2>Bản thảo luận · Đã và Sẽ</h2><p>Kết quả tháng {period}; kế hoạch tháng {draft.planPeriod}.</p><a aria-disabled={dirty} href={dirty?undefined:`/api/monthly/${period}/discussion.docx`}>Tải Word</a></div><div><h2>Bản nộp trường · A và B</h2><p>Báo cáo tháng {draft.planPeriod}: A.1, A.2, B.1, B.2 và kiến nghị. Mã đơn vị: 6.</p><a aria-disabled={dirty} href={dirty?undefined:`/api/monthly/${period}/school.docx`}>Tải Word</a></div></div>{dirty&&<p className="editor-hint">Lưu thay đổi trước khi tải Word.</p>}
- <p>{draft.records} báo cáo nguồn. {draft.counts.unknownPaperStatus>0&&`${draft.counts.unknownPaperStatus} công trình chưa có trạng thái xác định, không tính là đã được chấp nhận.`}</p>
- {Boolean(draft.reviewNotes?.length)&&<details className="product-settings"><summary>Nội dung cần bổ sung ({draft.reviewNotes?.length})</summary><p>Các mục này chỉ kèm bản thảo luận, chưa đưa vào bản báo cáo nộp trường. Sửa thông tin ở báo cáo nguồn rồi tổng hợp lại.</p>{draft.reviewNotes?.map((note,i)=><p key={i}>{note}</p>)}</details>}
- <div className="monthly-meta"><label>Tên kế hoạch chiến lược<input value={draft.strategyLabel} onChange={e=>{setDraft({...draft,strategyLabel:e.target.value});setDirty(true)}}/></label><label>Trưởng đơn vị<input value={draft.signatory} onChange={e=>{setDraft({...draft,signatory:e.target.value});setDirty(true)}}/></label></div>
- {Object.entries(labels).map(([key,label])=><label className="monthly-section" key={key}><strong>{label}</strong><textarea rows={Math.min(12,Math.max(3,(draft.sections[key]||'').split('\n').length+1))} value={draft.sections[key]||''} onChange={e=>{setDraft({...draft,sections:{...draft.sections,[key]:e.target.value}});setDirty(true)}}/></label>)}
+ {draft&&<>{draft.customized&&(!draft.assemblyVersion||draft.assemblyVersion<2)&&<p className="editor-hint">Đây là bản đã lưu bằng cách tổng hợp cũ. Bấm Tổng hợp lại để nối dòng và gộp mục trùng; kiểm tra nội dung trước khi lưu.</p>}<div className="monthly-downloads"><div><h2>Bản thảo luận · Đã và Sẽ</h2><p>Kết quả tháng {period}; kế hoạch tháng {draft.planPeriod}.</p><a aria-disabled={dirty} href={dirty?undefined:`/api/monthly/${period}/discussion.docx`}>Tải Word</a></div><div><h2>Bản nộp trường · A và B</h2><p>Báo cáo tháng {draft.planPeriod}: A.1 = A.2, B.1 = B.2. Mã đơn vị: 6.</p><a aria-disabled={dirty} href={dirty?undefined:`/api/monthly/${period}/school.docx`}>Tải Word</a></div></div>{dirty&&<p className="editor-hint">Lưu thay đổi trước khi tải Word.</p>}
+ <p>Nhập nội dung một lần cho mỗi phần. Khi xuất, A.1 được chép sang A.2 và B.1 được chép sang B.2.</p>
+ {Object.entries(labels).map(([key,label])=><label className="monthly-section" key={key}><strong>{label}</strong><textarea rows={Math.min(12,Math.max(3,(draft.sections[key]||'').split('\n').length+1))} value={draft.sections[key]||''} onChange={e=>{setDraft({...draft,sections:{...draft.sections,[key]:e.target.value,[key==='a1'?'a2':'b2']:e.target.value}});setDirty(true)}}/></label>)}
  </>}
  </section>
 }
'''

def digest(text):
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def apply_diff(source, lines):
    original = source.splitlines(keepends=True)
    result, cursor, i = [], 0, 0
    while i < len(lines):
        header = re.fullmatch(r"@@ -(\d+)(?:,\d+)? \+(\d+)(?:,\d+)? @@.*\n?", lines[i])
        if not header:
            raise ValueError("Invalid embedded diff header")
        start = max(0, int(header.group(1)) - 1)
        if start < cursor:
            raise ValueError("Overlapping diff")
        result.extend(original[cursor:start])
        cursor, i = start, i + 1
        while i < len(lines) and not lines[i].startswith("@@ "):
            line = lines[i]
            if line[0] in " -":
                if cursor >= len(original) or original[cursor] != line[1:]:
                    raise ValueError("Diff context does not match source")
                cursor += 1
            if line[0] in " +":
                result.append(line[1:])
            if line[0] not in " +-":
                raise ValueError("Invalid diff line")
            i += 1
    result.extend(original[cursor:])
    return "".join(result)


def read_text(path):
    return path.read_text(encoding="utf-8-sig") if path.exists() else ""


def main():
    cli = argparse.ArgumentParser(description="Update the MMLab monthly form after the automatic-receipt update.")
    cli.add_argument("--check", action="store_true", help="Check only; do not write files")
    cli.add_argument("--root", type=Path, default=Path(__file__).resolve().parent)
    args = cli.parse_args()
    root = args.root.resolve()
    if not (root / "run_local.py").is_file():
        sys.exit("Place this file beside run_local.py, or use --root PATH.")
    sections, key = {}, None
    lines = PATCH.splitlines(keepends=True)
    i = 0
    while i < len(lines):
        if lines[i].startswith("--- a/"):
            key = lines[i][6:].rstrip("\n")
            if i+1 >= len(lines) or lines[i+1] != "+++ b/" + key + "\n":
                sys.exit("Invalid patch file headers")
            sections[key] = []
            i += 2
        else:
            if key is None:sys.exit("Invalid embedded patch")
            sections[key].append(lines[i])
            i += 1
    if set(sections) != set(MANIFEST):sys.exit("Incomplete embedded patch")
    pending, incompatible = {}, []
    for rel, hashes in MANIFEST.items():
        path = (root / rel).resolve()
        if not path.is_relative_to(root):sys.exit("Source path escapes project")
        old = read_text(path)
        if path.exists() and digest(old) == hashes['after']:
            continue  # Already updated: no rewrite and no extra backup.
        matches = (not path.exists()) if hashes['before'] is None else (path.exists() and digest(old) == hashes['before'])
        if not matches:
            incompatible.append(rel)
            continue
        new = apply_diff(old, sections[rel])
        if digest(new) != hashes['after']:sys.exit("Patch checksum failed: " + rel)
        if rel.endswith('.py'):compile(new, rel, 'exec')
        pending[rel] = new
    if incompatible:
        sys.exit("STOP: source differs from the previous version. Nothing was written.\n" + "\n".join(incompatible))
    if not pending:
        print("Already updated. Rebuild frontend if needed, then restart the app.")
        return
    if args.check:
        print("Compatible: " + str(len(pending)) + " source files will be updated. Nothing written.")
        return
    stamp = datetime.now().strftime('%Y%m%d-%H%M%S-%f')
    backup = root / '_backups' / ('paired-ab-' + stamp)
    original = {}
    # Prepare all backups before replacing any source file.
    for rel in pending:
        path = root / rel
        original[rel] = path.read_bytes() if path.exists() else None
        if path.exists():
            saved = backup / rel
            saved.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, saved)
    backup.mkdir(parents=True, exist_ok=True)
    (backup / 'changes.json').write_text(json.dumps({'new_files':[r for r,b in original.items() if b is None], 'changed_files':list(pending)}, indent=2), encoding='utf-8')
    changed = []
    try:
        for rel, content in pending.items():
            path = root / rel
            path.parent.mkdir(parents=True, exist_ok=True)
            temp = None
            try:
                with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as f:
                    temp = Path(f.name)
                    f.write(content.encode('utf-8'))
                os.replace(temp, path)
                changed.append(rel)
            finally:
                if temp is not None and temp.exists():temp.unlink()
    except BaseException:
        for rel in reversed(changed):
            path = root / rel
            if original[rel] is None:path.unlink(missing_ok=True)
            else:path.write_bytes(original[rel])
        raise
    print("Updated " + str(len(pending)) + " source files.")
    print("Backup: " + str(backup))
    print("Next: npm ci --prefix frontend")
    print("      npm run build --prefix frontend")
    print("      python run_local.py")
    print("Restart the app using the same Python environment as before, then refresh the page.")


if __name__ == '__main__':
    main()
